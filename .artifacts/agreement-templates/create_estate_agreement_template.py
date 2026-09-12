from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("Estate_Agreement_Template.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "B8C2CC")


def set_table_width(table, width_dxa=9360):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def set_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, bold=False):
    paragraph = doc.add_paragraph(style=style)
    run = paragraph.add_run(text)
    set_font(run, bold=bold)
    return paragraph


def add_clause(doc, number, heading, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    heading_run = p.add_run(f"{number}. {heading}")
    set_font(heading_run, size=12, bold=True, color="1F4D78")
    body_p = add_para(doc, body)
    body_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body_p.paragraph_format.space_after = Pt(6)


def add_signature_block(doc, title, name_placeholder, date_placeholder):
    table = doc.add_table(rows=4, cols=2)
    table.autofit = False
    set_table_width(table)
    set_table_borders(table)
    widths = [Inches(2.1), Inches(4.4)]
    rows = [
        ("Name", name_placeholder),
        ("Signature", ""),
        ("Date", date_placeholder),
        ("Witness", "{{WitnessName}}"),
    ]
    title_cell = table.rows[0].cells[0].merge(table.rows[0].cells[1])
    mark_header_row(table.rows[0])
    title_cell.text = title
    set_cell_shading(title_cell, "F2F4F7")
    for paragraph in title_cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            set_font(run, bold=True, color="1F4D78")
    for index, (label, value) in enumerate(rows, start=1):
        if index >= len(table.rows):
            table.add_row()
        cells = table.rows[index].cells
        for width_index, width in enumerate(widths):
            cells[width_index].width = width
        cells[0].text = label
        cells[1].text = value
        for cell in cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_font(run, bold=(cell is cells[0]))
    doc.add_paragraph()


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = "Estate Agreement Template"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_font(run, size=9, color="6B7280")

    footer = section.footer.paragraphs[0]
    footer.text = "Prepared for DMS merge. Replace placeholders before execution."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        set_font(run, size=9, color="6B7280")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("PROPERTY LEASE / SALE AGREEMENT")
    set_font(run, size=18, bold=True, color="0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Agreement Reference: {{AgreementReference}}")
    set_font(run, size=11, bold=True, color="1F4D78")

    add_para(
        doc,
        "This template is intended for generation by the Document Management "
        "System. Keep the double-brace fields unchanged when uploading the "
        "template so the system can merge captured property, customer, lease, "
        "billing, and acquisition values.",
    )

    doc.add_heading("Agreement Summary", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    set_table_width(table)
    set_table_borders(table)
    headers = ["Field", "Value", "Field", "Value"]
    mark_header_row(table.rows[0])
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = text
        set_cell_shading(cell, "F2F4F7")
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_font(run, bold=True, color="1F4D78")
    rows = [
        ("Agreement date", "{{AgreementDate}}", "Move-in / start date", "{{MoveInDate}}"),
        ("Property / unit", "{{PropertyUnit}}", "Property reference", "{{PropertyReference}}"),
        ("Customer / lessee", "{{CustomerName}}", "Customer reference", "{{CustomerReference}}"),
        ("Request type", "{{RequestType}}", "Listing reference", "{{ListingReference}}"),
        ("Lease term", "{{LeaseTerm}}", "Billing start date", "{{BillingStartDate}}"),
        ("Currency", "{{Currency}}", "Rent / consideration", "{{PaymentAmount}}"),
    ]
    for row_values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_values):
            cells[index].text = value
            set_cell_margins(cells[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cells[index].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_font(run, bold=(index in (0, 2)))

    doc.add_heading("Parties", level=1)
    add_para(
        doc,
        "This Agreement is made between {{GrantorName}} of {{GrantorAddress}} "
        "('Owner/Lessor/Vendor') and {{GranteeName}} of {{GranteeAddress}} "
        "('Customer/Lessee/Purchaser').",
    )

    doc.add_heading("Property", level=1)
    add_para(
        doc,
        "The subject property is {{PropertyUnit}}, located at {{Location}}, "
        "with file or source reference {{SourceReference}}. The intended use "
        "is {{IntendedUse}}. Where applicable, the estimated size is "
        "{{EstimatedSize}} and the root of title is {{RootOfTitle}}.",
    )

    doc.add_heading("Commercial Terms", level=1)
    add_clause(
        doc,
        "1",
        "Term and commencement",
        "The tenancy, lease, sale, or acquisition arrangement starts on "
        "{{MoveInDate}} or such other date stated in the approved workflow. "
        "Billing shall start from {{BillingStartDate}}, which must not be "
        "earlier than the actual move-in or agreement start date.",
    )
    add_clause(
        doc,
        "2",
        "Payment",
        "The payment type is {{PaymentType}}. The agreed amount is "
        "{{PaymentAmount}} {{Currency}}. The payment schedule is "
        "{{PaymentSchedule}}.",
    )
    add_clause(
        doc,
        "3",
        "Deposit, rent, service charge, or ground rent",
        "The rent, deposit, service charge, ground rent, or other recurring "
        "charge shall be billed only after the signed agreement and the "
        "move-in / start date have been recorded in the system.",
    )

    doc.add_heading("Customer and Internal Approval", level=1)
    add_clause(
        doc,
        "4",
        "Customer approval",
        "The customer approval status is {{CustomerApprovalStatus}} and the "
        "approval date is {{CustomerApprovalDate}}.",
    )
    add_clause(
        doc,
        "5",
        "Internal approval",
        "Internal workflow approval shall be evidenced by {{WorkflowReference}} "
        "and the routed approval record. The current approval owner or role is "
        "{{NextApproverRole}} where applicable.",
    )
    add_clause(
        doc,
        "6",
        "Digital signature",
        "Where digital signing is used, the internal signer is "
        "{{InternalSignerName}}, signing as {{InternalSignerRole}}, with "
        "signature evidence reference {{InternalSignatureReference}}.",
    )

    doc.add_heading("Conditions", level=1)
    conditions = [
        "The customer shall comply with all property use rules, service obligations, estate regulations, and approved purpose restrictions.",
        "The owner or estate manager may withhold handover until all required signed documents, identity checks, and payment conditions are satisfied.",
        "Any variation, renewal, termination, assignment, or acquisition-stage change shall follow the applicable workflow approval route.",
        "Special conditions: {{SpecialConditions}}",
    ]
    for item in conditions:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_font(run)

    doc.add_heading("Document and Records Control", level=1)
    add_para(
        doc,
        "After execution, the signed agreement shall be uploaded or indexed "
        "under DMS reference {{AgreementDmsReference}} and property file "
        "reference {{PropertyFileReference}}. The generated agreement record is "
        "{{GeneratedAgreementReference}}.",
    )

    doc.add_heading("Execution", level=1)
    add_signature_block(doc, "Owner / Lessor / Vendor", "{{GrantorName}}", "{{GrantorSignatureDate}}")
    add_signature_block(doc, "Customer / Lessee / Purchaser", "{{CustomerName}}", "{{CustomerSignatureDate}}")
    add_signature_block(doc, "Internal Approval / Estate Manager", "{{InternalSignerName}}", "{{InternalSignatureDate}}")

    doc.add_heading("Placeholder Field Checklist", level=1)
    add_para(
        doc,
        "Core fields used by this template: {{AgreementReference}}, "
        "{{AgreementDate}}, {{PropertyUnit}}, {{PropertyReference}}, "
        "{{ListingReference}}, {{CustomerName}}, {{CustomerReference}}, "
        "{{RequestType}}, {{MoveInDate}}, {{BillingStartDate}}, {{LeaseTerm}}, "
        "{{PaymentType}}, {{PaymentAmount}}, {{PaymentSchedule}}, {{Currency}}, "
        "{{GrantorName}}, {{GrantorAddress}}, {{GranteeName}}, "
        "{{GranteeAddress}}, {{Location}}, {{IntendedUse}}, {{EstimatedSize}}, "
        "{{RootOfTitle}}, {{WorkflowReference}}, {{NextApproverRole}}, "
        "{{InternalSignerName}}, {{InternalSignerRole}}, "
        "{{InternalSignatureReference}}, {{AgreementDmsReference}}, "
        "{{PropertyFileReference}}, {{GeneratedAgreementReference}}, and "
        "{{SpecialConditions}}.",
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
    print(OUT)
