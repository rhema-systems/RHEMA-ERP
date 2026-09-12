from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("TDC_Property_Sale_Agreement_Template.docx")


BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(33, 37, 41)
MUTED = RGBColor(90, 99, 110)
LIGHT_FILL = "F2F4F7"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.append(tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)
    for row in table.rows:
        for index, width in enumerate(widths):
            cell = row.cells[index]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_run(paragraph, text, bold=False, size=None, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    return run


def add_paragraph(doc, text="", style=None, bold=False):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    if text:
        add_run(paragraph, text, bold=bold, color=DARK)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.keep_with_next = True
    add_run(paragraph, text)
    return paragraph


def add_clause(doc, title, body):
    add_heading(doc, title, 2)
    paragraph = add_paragraph(doc, body)
    paragraph.paragraph_format.keep_together = False


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.add_run(item)


def add_detail_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0]
    hdr.cells[0].text = "Agreement item"
    hdr.cells[1].text = "Value"
    set_repeat_table_header(hdr)
    for cell in hdr.cells:
        set_cell_shading(cell, LIGHT_FILL)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = DARK
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_table_width(table, [3000, 6360])
    return table


def add_signature_table(doc):
    add_heading(doc, "Execution", 1)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Seller / authorised officer"
    hdr[1].text = "Purchaser"
    for cell in hdr:
        set_cell_shading(cell, LIGHT_FILL)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    rows = [
        ("Name: ______________________________", "Name: {{ApplicantName}}"),
        ("Title: ______________________________", "Customer reference: {{CustomerReference}}"),
        ("Signature: ___________________________", "Signature: ___________________________"),
        ("Date: _______________________________", "Date: _______________________________"),
        ("Witness: ____________________________", "Witness: ____________________________"),
    ]
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
    set_table_width(table, [4680, 4680])


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, RGBColor(31, 77, 120), 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def build():
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_run(header, "Estate / Property Management", size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(footer, "Template code: EST-SALE-AGREEMENT | Source: Estate / Property Management -> Central DMS", size=8, color=MUTED)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    add_run(title, "PROPERTY SALE AGREEMENT", bold=True, size=18, color=BLUE)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    add_run(subtitle, "Generated from Estate Property Management listing application {{CaseReference}}", size=10, color=MUTED)

    add_detail_table(
        doc,
        [
            ("Agreement date", "{{AgreementDate}}"),
            ("Estate / workflow reference", "{{CaseReference}}"),
            ("Purchaser", "{{ApplicantName}}"),
            ("Customer reference", "{{CustomerReference}}"),
            ("Property / unit", "{{PropertyNumber}}"),
            ("Property reference", "{{PropertyReference}}"),
            ("Listing reference", "{{ListingReference}}"),
            ("Request type", "{{RequestType}}"),
            ("Agreed purchase price", "{{Currency}} {{PaymentAmount}}"),
            ("Prepared by", "{{PreparedBy}}"),
        ],
    )

    add_heading(doc, "Parties", 1)
    add_paragraph(
        doc,
        "This Property Sale Agreement is made between the seller or authorised estate owner represented by the Estate / Property Management function (the Seller) and {{ApplicantName}} with customer reference {{CustomerReference}} (the Purchaser).",
    )
    add_paragraph(
        doc,
        "The parties record that the Purchaser submitted a purchase bid under Estate reference {{CaseReference}} and that the Seller has approved the sale subject to this agreement, internal controls, customer acceptance, execution, payment, and completion of conveyance or registration requirements.",
    )

    add_clause(
        doc,
        "Property",
        "The property sold under this agreement is {{PropertyNumber}}, referenced as {{PropertyReference}} and listed under {{ListingReference}}. The property is sold together with the rights and appurtenances that the Seller is legally able to transfer, subject to estate records, approved plans, reservations, covenants, encumbrances, and applicable statutory requirements.",
    )
    add_clause(
        doc,
        "Purchase price",
        "The agreed purchase price is {{Currency}} {{PaymentAmount}}. This amount records the approved commercial decision for the purchase bid and excludes any tax, stamp duty, registration fee, statutory charge, service charge, utility deposit, legal cost, or administrative charge unless expressly stated in writing by the Seller.",
    )
    add_clause(
        doc,
        "Payment and settlement",
        "The Purchaser shall settle the purchase price and any applicable charges by the deadline and payment channel communicated by the Seller. The Seller may withhold possession, conveyance, title transfer, or registration steps until cleared funds and required evidence of payment have been confirmed.",
    )
    add_clause(
        doc,
        "Conditions precedent",
        "Completion of this agreement is conditional on successful internal approval, Legal review, customer acceptance, execution by the parties, verification of submitted documents, payment confirmation, and any additional statutory or estate governance requirements that apply to the property.",
    )
    add_bullets(
        doc,
        [
            "The Purchaser shall provide accurate identity, eligibility, and payment information.",
            "The Seller may request additional evidence where required for due diligence or records control.",
            "The Seller may suspend completion if a competing legal claim, restriction, or material discrepancy is identified.",
        ],
    )
    add_clause(
        doc,
        "Possession and handover",
        "Possession shall pass only after the Seller confirms that all completion conditions have been met and a written handover or possession instruction has been issued. Any proposed move-in, handover, inspection, or key release date remains provisional until approved in the workflow.",
    )
    add_clause(
        doc,
        "Purchaser undertakings",
        "The Purchaser undertakes to comply with estate rules, planning and building controls, use restrictions, payment obligations, and all covenants attached to the property. The Purchaser shall not assign, occupy, alter, mortgage, lease, or otherwise deal with the property before completion except with written approval from the Seller.",
    )
    add_clause(
        doc,
        "Seller undertakings",
        "The Seller shall process the approved transaction through the required internal, legal, document-control, and records steps. The Seller shall provide reasonable completion documentation after the Purchaser has satisfied the conditions precedent.",
    )
    add_clause(
        doc,
        "Default",
        "If the Purchaser fails to complete payment, provide required documents, execute the agreement, or satisfy completion conditions within the required period, the Seller may cancel the approval, reopen the property for allocation or sale, retain permissible administrative charges, or take any other action allowed by policy and law.",
    )
    add_clause(
        doc,
        "Records and notices",
        "The Estate / Property Management workflow, Central DMS record, and related legal or finance records form part of the administrative record for this transaction. Notices may be issued through the customer portal, email, written letter, or any contact address supplied by the Purchaser.",
    )
    add_clause(
        doc,
        "Governing law and review",
        "This agreement is subject to applicable law and final Legal review. Any correction required by Legal, Finance, Records, or authorised management must be resolved before the agreement is treated as fully executed.",
    )

    add_heading(doc, "Completion Checklist", 1)
    checklist = doc.add_table(rows=1, cols=3)
    checklist.style = "Table Grid"
    checklist.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = checklist.rows[0]
    hdr.cells[0].text = "Control"
    hdr.cells[1].text = "Owner"
    hdr.cells[2].text = "Status / reference"
    set_repeat_table_header(hdr)
    for cell in hdr.cells:
        set_cell_shading(cell, LIGHT_FILL)
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in [
        ("Management approval", "Property Manager", "{{CaseReference}}"),
        ("Legal review", "Legal", "Pending / ____________________"),
        ("Customer acceptance", "Customer", "{{CustomerApprovalStatus}}"),
        ("Payment confirmation", "Finance AR", "{{Currency}} {{PaymentAmount}}"),
        ("Final signed agreement", "Central DMS", "{{GeneratedAgreementReference}}"),
    ]:
        cells = checklist.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    set_table_width(checklist, [3400, 2400, 3560])

    add_signature_table(doc)

    add_heading(doc, "Administrative Notes", 1)
    add_paragraph(doc, "Generated on {{Today}} from workflow reference {{WorkflowReference}}.")
    add_paragraph(doc, "Next approver / responsible role at generation: {{NextApproverRole}}.")
    add_paragraph(doc, "This template is intended for workflow-driven generation and should be reviewed by Legal before external execution.")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
