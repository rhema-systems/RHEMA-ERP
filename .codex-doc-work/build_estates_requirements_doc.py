from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\Users\etwen\Documents\GitHub\RHEMA-ERP\docs\TDC_Estates_Module_Requirements_and_Gap_Analysis.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
BORDER = "D9E2F3"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        tbl.insert(1, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table, color="BFBFBF"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def style_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.text = "TDC Estates ERP Requirements and Gap Analysis"
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(85, 85, 85)
    add_page_number(footer)


def add_title(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("TDC Estates Module Requirements and Gap Analysis")
    r.font.name = "Calibri"
    r._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    r._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(INK)

    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(12)
    run = sub.add_run("Prepared from TDC Estates Department workshop questionnaire responses and the Estates Operational Manual, 2026.")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(85, 85, 85)

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.style = "Table Grid"
    rows = [
        ("Document purpose", "Module-owner requirements document for ERP solution scoping, validation, and implementation planning."),
        ("Department/module", "TDC Estates: land/property/estate management and home ownership scheme operations."),
        ("Source documents", "TDC Estates Questionnaire Response - Final ERP Submission, 26 June 2026; Estates Operational Manual - Final ERP Submission, 26 June 2026."),
        ("Prepared date", "14 July 2026"),
    ]
    for row, (label, value) in zip(meta.rows, rows):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT_FILL)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)
            set_cell_margins(cell)
    set_table_geometry(meta, [2300, 7060])
    set_table_borders(meta)


def add_callout(doc, title, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    r.font.size = Pt(11)
    p.add_run("\n" + text)
    set_table_geometry(table, [9360])
    set_table_borders(table, color=BORDER)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_table(doc, headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], LIGHT_FILL)
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor.from_string(INK)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(9)
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    doc.add_paragraph()
    return table


def build_doc():
    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "This document converts the TDC Estates Department workshop response and operational manual into ERP module requirements. "
        "It is intended for module-owner validation before solution design and configuration."
    )
    add_callout(
        doc,
        "Overall assessment",
        "Core business processes are well documented procedurally, but many are manual, register-based, and dependent on file movement. "
        "The ERP should therefore prioritize a controlled property and customer record, workflow routing, document tracking, Revenue/Finance integration, "
        "case and work-order management, reporting, and system-level audit trails.",
    )

    add_table(
        doc,
        ["Area", "Current position", "ERP implication"],
        [
            ("Records and files", "Estate registers, rent cards, rent registers, ledger cards, and physical property files are maintained by Records and sections.", "Create digital master records, document repositories, and file movement tracking with audit history."),
            ("Allocations and leases", "Applications, proposal letters, offer letters, Right of Entry letters, cadastral plans, lease requests, Legal processing, Lands Commission registration, and record updates are established.", "Configure end-to-end workflows, approvals, templates, status milestones, and prerequisite checks."),
            ("Property and housing", "About 6,000 rental and home ownership scheme units are managed, including rental compliance, conversion to HOS, rent cards, rent rolls, and offer letters.", "Track unit inventory, tenancy status, rent position, HOS conversion, sale progress, and ownership/title milestones."),
            ("Facilities and estate services", "Facilities Management manages TDC Towers, Site 3, Community 26, sanitation, landscaping, security, contractors, utilities, and debt recovery.", "Add facilities asset register, work orders, service provider tasks, incident logs, utility/service charge tracking, and contractor performance reports."),
            ("Controls", "Approvals and signatories are defined by role, but segregation of duties, system audit logs, policy enforcement, and exception registers are not formalized in system form.", "Configure approval matrix, role permissions, audit logs, exception/override workflows, and control reports."),
        ],
        [1800, 3800, 3760],
    )

    doc.add_heading("2. Scope", level=1)
    doc.add_paragraph("The Estates ERP module should cover these operating areas:")
    add_bullets(
        doc,
        [
            "Land administration records that touch Estates workflows, including allocation, transfer, assignment, mortgage consent, lease preparation, lease renewal, regularization, and traditional land documentation.",
            "Property management and housing operations, including rental housing, home ownership scheme units, serviced plots, partially serviced plots, commercial/facilities managed properties, and customer records.",
            "Estate and facilities management, including sanitation, landscaping, security, utilities, common-area upkeep, contractor coordination, inspections, service charges, and incidents.",
            "Customer service, case tracking, correspondence, notices, complaints, dispute history, and file/letter movement.",
            "Reporting, approvals, controls, audit trail, and integrations with Revenue, Finance, Legal, Development, Marketing, Internal Audit, and ICT.",
        ],
    )

    doc.add_heading("3. Source-Based Current State", level=1)
    add_table(
        doc,
        ["Process domain", "Confirmed current-state process"],
        [
            ("Organization", "The Estates Department operates through Lands/Partially Serviced, Traditional, Housing, Serviced Plots, Regularization, Facilities Management, Records, and Secretarial/Registry functions under the Head of Estates."),
            ("Records", "Estate registers capture allocations, right of entry/date of tenancy, lease term, ground rent, acreage or plot size, lessee details, and property file references. HOS ledger cards and rent registers are also maintained."),
            ("Application intake", "Forms are purchased after payment and receipt, then logged by form type, applicant name, property number, receipt number, and authorizing officer."),
            ("Allocation", "Applications are compiled and referred for consideration; proposal letters communicate plot size, location, amount payable, and payment deadline. MD approval/signature is required for key letters and allocations."),
            ("Serviced plots and HOS", "Serviced plots are managed across Communities 14, 18, 20, 22, 24, and 25. HOS units are managed in identified communities and Site 3 blocks. Offer letters and Right of Entry letters are prepared after payment milestones."),
            ("Housing", "The Housing Section manages about 6,000 rental and HOS units, rent cards, rent registers, inspections, rent rolls, quarterly reports, rental sale analysis, transfers, mortgage consents, and dispute support."),
            ("Revenue dependency", "Most processes require Revenue confirmation of payment or arrears status. Revenue and Estate Records are both amended after approved changes."),
            ("Legal and title", "Lease, assignment, deed variation, consent to mortgage, and registration processes require Legal Department action and Lands Commission milestones."),
            ("Estate services", "Facilities Management manages common services, contractors, utilities, TDC Towers, Site 3 properties, Community 26, customer complaints, billing support, and debt recovery with Revenue."),
            ("Reporting", "Weekly, monthly, and quarterly reports exist for officers/assistants and section activity; however, several management and estate service report templates are not yet defined."),
        ],
        [2300, 7060],
    )

    doc.add_heading("4. Functional Requirements", level=1)

    requirements = [
        (
            "FR-01",
            "Digital estate/property master register",
            "The system shall maintain a single searchable register for plots, lands, rental units, HOS units, commercial lettable spaces, common assets, blocks, sites, and estates.",
            "Must include property number, location/community, land use, plot/unit size, acreage, lease term, date of tenancy/right of entry, ground rent, current status, linked customer, and file reference.",
            "High",
        ),
        (
            "FR-02",
            "Customer and occupant records",
            "The system shall maintain tenant, lessee, applicant, purchaser, occupant, joint tenant, assignee, mortgagee, and contact records with full relationship history.",
            "Support address changes, additions of names, transfers, assignments, mortgages, and multiple tenants for a unit.",
            "High",
        ),
        (
            "FR-03",
            "Application and form intake",
            "The system shall capture purchase and submission of Home Ownership Scheme Forms, Rental Unit Forms, Estate Transfer Forms, regularization forms, and land/plot applications.",
            "Capture receipt number, purchase date, applicant name, property number, application type, documents received, receiving officer, and current owner/team.",
            "High",
        ),
        (
            "FR-04",
            "File and correspondence tracking",
            "The system shall replace manual incoming/outgoing books with electronic file, letter, memo, and form movement tracking.",
            "Track sender, receiver, date received, action/comment, officer assigned, due date, dispatch status, and location history.",
            "High",
        ),
        (
            "FR-05",
            "Allocation workflow",
            "The system shall support land, serviced plot, HOS unit, and rental unit allocation workflows from application to approval, proposal letter, payment, offer letter, and record update.",
            "Include availability checks, PAC/committee routing where applicable, MD approval/signature steps, deposit/payment prerequisites, and applicant notification.",
            "High",
        ),
        (
            "FR-06",
            "Proposal, offer, and Right of Entry letters",
            "The system shall generate controlled letter templates for proposal letters, offer letters, Right of Entry letters, rate revision notices, demand letters, acknowledgements, and completion letters.",
            "Templates should pull approved property, customer, fee, receipt, lease term, ground rent, and date fields from system records.",
            "High",
        ),
        (
            "FR-07",
            "Billing prerequisite and arrears controls",
            "The system shall enforce arrears clearance before configured processes proceed.",
            "Processes include searches, certified true copies, transfers, assignments, leases, recognition of tenancy, conversion to HOS, and other workflows defined by policy.",
            "High",
        ),
        (
            "FR-08",
            "Revenue/Finance integration",
            "The system shall integrate or reconcile with Revenue/Finance for ground rent, land management fees, regularization fees, proposal payments, deposits, common area maintenance, rent, service charges, receipts, and arrears balances.",
            "Receipting is handled outside Estates today, so ERP design must clearly define data ownership and synchronization.",
            "High",
        ),
        (
            "FR-09",
            "Lease and title milestone tracking",
            "The system shall track cadastral plan request, Legal lease request, invoice/demand letter, fee payment, Lands Commission registration, legal detachment, and Revenue/Estate Records amendment milestones.",
            "Include lease terms, renewal eligibility, surrender and renewal status, deed of variation, assignment, mortgage consent, and title document links.",
            "High",
        ),
        (
            "FR-10",
            "Housing and HOS lifecycle",
            "The system shall support rental tenancy recognition, transfer, rent cards, rent registers, rental-to-HOS conversion, installment sale, full payment confirmation, HOS offer letter, lease request, and completion status.",
            "Must handle multi-tenanted units where all tenants must complete payment before HOS offer letters are issued.",
            "High",
        ),
        (
            "FR-11",
            "Inspection management",
            "The system shall record site inspections, house inspections, substantial development inspections, compliance observations, tenancy compliance checks, encroachment referrals, and weekly inspection reports.",
            "Inspection records should include checklist, officer, date, findings, photos/documents, next action, referral, and close-out status.",
            "High",
        ),
        (
            "FR-12",
            "Facilities work orders",
            "The system shall create, approve, assign, track, and close maintenance/facilities work orders for TDC-managed estates and facilities.",
            "Include request source, priority, asset/location, contractor, approval threshold, estimated/actual cost, completion evidence, and feedback.",
            "High",
        ),
        (
            "FR-13",
            "Contractor and service provider oversight",
            "The system shall track outsourced cleaners, gardeners, security, external contractors, service contracts, assignments, SLA targets, performance issues, and contractor reports.",
            "Should support task schedules, inspections, non-performance notices, and service-level dashboards.",
            "Medium",
        ),
        (
            "FR-14",
            "Estate service charge and utility management",
            "The system shall support Common Area Maintenance charges, service charges, utility pass-throughs, shared-cost recoveries, billing disputes, and debt recovery actions.",
            "Rates, formulas, billing cycles, dispute rules, and responsible approvers must be configurable once finalized.",
            "High",
        ),
        (
            "FR-15",
            "Customer service and case management",
            "The system shall log inquiries, complaints, notices, dispute cases, resident communications, call-centre interactions, and escalation history across land, rental, HOS, and facilities customers.",
            "Include case category, property/customer link, SLA, owner, escalation, communication log, documents, and resolution.",
            "High",
        ),
        (
            "FR-16",
            "Document repository",
            "The system shall store and retrieve application forms, statutory declarations, site plans, receipts, offer letters, rent cards, leases, assignments, mortgage consents, court documents, reports, notices, and approval forms.",
            "Support document checklists by process and role-based access.",
            "High",
        ),
        (
            "FR-17",
            "Reporting and dashboards",
            "The system shall produce operational, revenue, arrears, allocation, lease, inspection, facilities, contractor, incident, and management reports.",
            "At minimum, include quarterly reports already prepared by Lands, Serviced Plots, Traditional, Housing, HOS, and Facilities teams.",
            "High",
        ),
        (
            "FR-18",
            "Approvals, controls, and audit trails",
            "The system shall enforce role-based approvals, signatory rules, segregation of duties, exception routing, and immutable audit trails for record changes and approvals.",
            "Must capture who changed what, when, old/new values, approval decision, reason, and supporting document.",
            "High",
        ),
    ]
    add_table(
        doc,
        ["ID", "Requirement", "System statement", "Notes / acceptance considerations", "Priority"],
        requirements,
        [800, 1800, 3300, 2600, 860],
    )

    doc.add_heading("5. Business Rules and Approval Requirements", level=1)
    rules = [
        ("BR-01", "Most departmental letters, search reports, certified true copies, amendments, lease acknowledgements, and approval forms are signed by the Head of Estates."),
        ("BR-02", "The Managing Director approves or signs key transfers, allocations, additional land approvals, offer letters, rent cards, fee/demand letters, assignment fee letters, layout revision letters, and lease renewal invoices."),
        ("BR-03", "Ground rent arrears must be cleared before many processes proceed, including search, certified true copy, transfer, assignment, lease preparation, tenancy recognition, and conversion to HOS."),
        ("BR-04", "Records must be amended in both Revenue and Estate Records after approved changes such as transfers, assignments, mortgages, leases, address changes, and HOS conversion."),
        ("BR-05", "Cadastral plans required for registration must comply with LI 1444."),
        ("BR-06", "Serviced plots require deposit/reservation and payment milestones before offer and right-of-entry documentation. The questionnaire indicates 40 percent deposit for serviced plots, while the manual describes full payment before offer/right-of-entry steps; this requires confirmation."),
        ("BR-07", "Home ownership scheme units use deposits and payment schedules handled with Marketing input. The questionnaire states 50 percent deposit for HOS units; the operational manual describes offer preparation after payment completion; this requires process confirmation."),
        ("BR-08", "For purchase of rental units, available models include full payment within one year, lump sum, or 70 percent within one year plus 30 percent over two years at Treasury bill rate plus two percentage points; default may revert the house to rental and offset payments against prevailing rent."),
        ("BR-09", "Lease renewal with unexpired term of ten years and below is handled directly; more than ten years is referred to the Lease Renewal Technical Committee."),
        ("BR-10", "Transfer and consent fees are set at 10 percent of the applicable Land Management Fee, value, or selling price, subject to confirmed fee policy."),
        ("BR-11", "Traditional Council allocations may be rejected where there is prior allocation, dispute, unclear signatories, or non-compliance with TDC documentation requirements."),
        ("BR-12", "Change of land use requests are referred to the relevant company committee before the applicant is informed of the outcome."),
    ]
    add_table(doc, ["Rule ID", "Rule / control requirement"], rules, [1200, 8160])

    doc.add_heading("6. Integration Requirements", level=1)
    add_table(
        doc,
        ["Interface", "Required exchange", "Reason"],
        [
            ("Revenue / Finance", "Receipts, arrears balances, rent, ground rent, fees, invoices, service charges, payment confirmations, debt recovery status.", "Estates processes depend on payment and arrears confirmation; Revenue performs receipting."),
            ("Legal", "Lease requests, assignments, mortgage consents, deed variations, legal correspondence, registration status, detachment evidence.", "Legal prepares and manages key title/contract documents."),
            ("Development / Building Inspectorate", "Building permit verification, cadastral/site plan request and completion, development compliance, change of use routing.", "Several Estates workflows require Development confirmation before approval/signature."),
            ("Marketing", "HOS and serviced plot referrals, deposit confirmation, payment plans, completion letters, customer communication for sales schemes.", "Marketing currently initiates or controls parts of HOS and serviced plot sales processes."),
            ("Internal Audit", "Approval logs, change history, exception registers, segregation-of-duties exceptions, report packs.", "Controls are currently procedural and require system evidence."),
            ("ICT / Document management", "Scanning, repository access, retention, permissions, backups, report exports, workflow notifications.", "Physical files and manual registers are core operational pain points."),
        ],
        [1900, 4300, 3160],
    )

    doc.add_heading("7. Reporting Requirements", level=1)
    add_table(
        doc,
        ["Report group", "Minimum reports"],
        [
            ("Land and title", "Applications received/processed, transfers and expected fees, proposal letters and expected revenue, lease applications, assignment applications, building permit verifications, offer/right-of-entry letters, mortgage applications, lease renewals."),
            ("Property and tenancy", "Rent roll, rent arrears, rent card status, occupancy status, tenancy recognition, transfers, HOS conversions, notices/reminders, court/dispute support."),
            ("Serviced plots and HOS", "Payments received, cumulative project payments from inception, debtor lists, deposits, installment schedules, completion letters, offer letters, lease processing, pending customer actions."),
            ("Facilities and estate services", "Work orders, open/closed maintenance tasks, contractor performance, inspections, incidents/service failures, utilities, service charges, debt recovery, parking/security/waste/landscaping activities."),
            ("Management and Board", "Monthly and quarterly performance, revenue, arrears, allocation throughput, lease/title backlog, risk/issues, complaints, service delivery, and control exceptions."),
            ("Audit/control", "User activity, record changes, approval cycle time, overrides, rejected applications, prerequisite bypasses, segregation-of-duties exceptions, missing documents."),
        ],
        [2400, 6960],
    )

    doc.add_heading("8. Gap Register", level=1)
    gaps = [
        ("G-01", "Manual file and register dependency", "Physical files, incoming/outgoing books, ledger cards, rent cards, and registers remain core records.", "High", "Digitize master records, document repository, file movement tracking, and searchable audit history."),
        ("G-02", "Storage and retrieval constraint", "Cabinet Room capacity and manual retrieval are identified challenges.", "High", "Create document scanning/indexing backlog plan and define retention/archive rules."),
        ("G-03", "No formal work-order lifecycle", "Facilities maintenance is coordinated, but no logged request-to-completion work-order process is in place.", "High", "Implement facilities work order, approvals, assignment, evidence, feedback, and close-out."),
        ("G-04", "No structured handover/return and defects process", "Inspections occur, but structured handover/return procedures and defect close-out are not documented.", "High", "Define handover, return, snag/defect, post-sale support, and customer acknowledgement workflow."),
        ("G-05", "No formal customer/call-centre case system", "Inquiries and complaints are handled by Secretariat/Facilities officers, but no formal case logging/escalation platform exists.", "High", "Implement case management with SLA, categories, communication log, escalation, and resolution evidence."),
        ("G-06", "Estate charge and shared-cost formula not standardized", "CAM exists for TDC Towers, but a documented formula across estates and dispute process is not established.", "High", "Confirm service charge formulas, billing cycles, dispute approval, and integration with Revenue."),
        ("G-07", "Common asset register missing", "Meters, streetlights, open spaces, incidents, and service failures are not maintained in a structured asset/service register.", "Medium", "Create common asset master data for facilities, utilities, open spaces, parking, lighting, drainage, and related inspections."),
        ("G-08", "Incident and service failure escalation missing", "A formal incident and service failure logging/escalation system is not in place.", "High", "Configure incident categories, severity, escalation, notifications, investigation, and close-out."),
        ("G-09", "Resident/community communication workflow missing", "Facilities handles complaints and communication, but structured resident notices, approvals, and community escalation are not documented.", "Medium", "Define resident communication workflow, notice templates, community issue register, and escalation rules."),
        ("G-10", "Eligibility scoring and affordability assessment incomplete", "Formal affordability assessments, employer/public-worker categories, and structured eligibility scoring are not established.", "Medium", "Confirm whether future HOS products require scoring, ranking, qualification documents, and approval transparency."),
        ("G-11", "Post-sale defect support unclear", "Structured post-sale defect reporting and customer support beyond general client service is not documented.", "High", "Define post-sale warranty/defect policy, support categories, timelines, and responsible teams."),
        ("G-12", "Estate management report templates undefined", "Contractor performance, incidents, utilities, and estate service report templates are not established.", "Medium", "Design standard facilities and estate services reports before build/configuration."),
        ("G-13", "Management and Board report pack undefined", "Quarterly reporting exists, but a defined Board pack covering performance, revenue, arrears, service delivery, and risk is not established.", "Medium", "Agree management dashboard and Board report pack layout, KPIs, cadence, and data owners."),
        ("G-14", "Segregation-of-duties matrix missing", "Some practical separation exists, but no formal role-combination prevention matrix is documented.", "High", "Define incompatible roles and approval restrictions for records, billing, allocation, maintenance, and receipting."),
        ("G-15", "System-level audit log missing", "File and letter movement is tracked manually; system-level record change and approval audit is a future-state need.", "High", "Implement immutable audit logs for record changes, approvals, letters, documents, and payment-status dependencies."),
        ("G-16", "Policy enforcement is procedural", "Rules such as arrears clearance, fee calculation, lease terms, and LI 1444 requirements are enforced by officers rather than system controls.", "High", "Translate policy rules into configurable validations, prerequisites, and exception workflows."),
        ("G-17", "Exception/override register missing", "Discretionary decisions exist, but no formal escalation and override approval register is established.", "High", "Configure exception request, reason, approver, expiry, audit trail, and reporting."),
        ("G-18", "New land acquisition process outside Tema not fully documented", "The department advises on acquisition options, but a complete new land acquisition workflow is not documented.", "Medium", "Run follow-up workshop with Land Administration, Legal, Survey/Technical, Finance, and Executive Management."),
        ("G-19", "Compensation/community engagement framework incomplete", "Valuations and dispute support exist, but structured compensation and community engagement for new acquisition is not documented.", "Medium", "Define compensation decision, stakeholder engagement, approvals, and communication requirements."),
        ("G-20", "Facilities SOP detail incomplete in operational manual", "The manual identifies Facilities & Maintenance structurally, while detailed facilities process data is mostly in questionnaire responses.", "Medium", "Validate facilities scope with Facilities Management and update SOPs or implementation notes before build."),
        ("G-21", "Street lighting and drainage service ownership unclear", "Questionnaire notes security, sanitation, landscaping, and common area upkeep; street lighting and drainage are not currently described as managed services.", "Low", "Confirm whether these services are in scope, outsourced, handled by another department, or excluded."),
    ]
    add_table(
        doc,
        ["Gap ID", "Gap", "Evidence / current limitation", "Priority", "Recommended action"],
        gaps,
        [800, 1900, 3100, 800, 2760],
    )

    doc.add_heading("9. Open Clarifications for Follow-Up", level=1)
    add_numbered(
        doc,
        [
            "Confirm whether this requirements document is for a combined Estates module or should be split into Land Management, Property Management, Estate/Facilities Management, and Home Ownership Scheme sub-documents.",
            "Confirm deposit rules and payment prerequisites for serviced plots and HOS units, because the questionnaire and operational manual describe different payment triggers.",
            "Confirm the exact approval path for Property Allocation Committee decisions, Managing Director approvals, Lease Renewal Technical Committee decisions, Change of Land Use Committee decisions, and emergency maintenance approvals.",
            "Confirm which systems currently hold Revenue/Finance receipts, arrears, invoices, rent, ground rent, service charge, and customer payment data.",
            "Confirm whether Marketing owns HOS sales payment plans and completion letters or whether the ERP workflow should route this through Estates with Marketing approval.",
            "Confirm whether TDC wants customer self-service or staff-only case management for inquiries, complaints, notices, payment plan visibility, and defect reporting.",
            "Confirm reporting cadence and audience for Board, management, Head of Estates, Estate Managers, Officers, Internal Audit, and Finance.",
            "Confirm which historical files should be migrated first: active tenants/lessees, arrears cases, active applications, HOS debtors, pending leases, litigation/disputes, or all records.",
        ],
    )

    doc.add_heading("10. Implementation Priority", level=1)
    add_table(
        doc,
        ["Phase", "Recommended scope", "Reason"],
        [
            ("Phase 1 - Foundation", "Property/customer master data, digital file tracking, document repository, Revenue/Finance payment status integration, core approvals, audit trail.", "These are prerequisites for reliable Estates workflows and reduce the biggest manual-record risks."),
            ("Phase 2 - Core operations", "Allocation, proposal/offer/right-of-entry letters, lease/title milestones, housing/HOS lifecycle, rent register/rent roll, arrears controls, standard reports.", "These cover the highest-volume Estates operational processes."),
            ("Phase 3 - Facilities and service delivery", "Work orders, common asset register, inspections, contractor tasks, service charges, utilities, incidents, resident cases, contractor performance reports.", "These address the most explicit facilities and estate-service gaps."),
            ("Phase 4 - Advanced controls and analytics", "Board dashboards, SLA analytics, exception analytics, document retention automation, customer portal, mobile inspections, performance KPIs.", "These strengthen governance and management visibility after core data is stable."),
        ],
        [1700, 5100, 2560],
    )

    doc.add_heading("11. Acceptance Criteria for Module-Owner Sign-Off", level=1)
    add_bullets(
        doc,
        [
            "Every requirement has a named business owner and confirming department.",
            "Every gap has a decision: implement now, implement later, document only, or exclude from project scope.",
            "Approval matrix and signatory rules are confirmed by Estates, Legal, Finance, and Executive Management.",
            "Revenue/Finance integration responsibility and source of truth are agreed.",
            "Data migration scope and document digitization approach are approved.",
            "Reports and dashboards are signed off with sample templates or confirmed field lists.",
            "Open clarifications are resolved or formally accepted as project risks before development/configuration begins.",
        ],
    )

    doc.add_heading("12. Source Traceability", level=1)
    add_table(
        doc,
        ["Source document", "Information used"],
        [
            ("TDC Estates Questionnaire Response - Final ERP Submission, 26 June 2026", "Organization/governance response, property management questionnaire, estate management questionnaire, home ownership scheme questionnaire, reporting, controls, approval matrix, appendix interview prompts, and documented gaps."),
            ("Estates Operational Manual - Final ERP Submission, 26 June 2026", "Department structure, secretarial/registry duties, records section, land and lease workflows, serviced plots, housing schedule, HOS conversion, offer/right-of-entry preparation, quarterly reporting, regularization, and manual controls."),
        ],
        [3500, 5860],
    )

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
